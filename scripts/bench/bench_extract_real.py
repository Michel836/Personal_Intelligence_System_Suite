# ruff: noqa
"""Representative extraction benchmark for the production ExtractionManager.

Unlike the tiny-fixture benchmark, this generates non-trivial documents for
every production format (PDF/DOCX/XLSX/ODT/ODS/TXT/HTML/JSON) and runs the two
real paths:

* ``sequential``: the loop used by scripts/extract_content.py (extract_single);
* ``batch_w``: ExtractionManager.extract_batch with w thread-pool workers
  (the AutoExtractor / archive path).

Correctness is compared across worker counts (success count and characters).
"""
from __future__ import annotations

import argparse
import json
import sys
import time
from pathlib import Path

from loguru import logger

logger.remove()

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from scripts.bench.common import ResourceSampler, json_print, load_conditions, wipe  # noqa: E402

LOREM = (
    "Le present contrat de service definit les obligations des parties, la "
    "confidentialite des informations echangees, le calendrier de livraison, les "
    "conditions financieres et la clause de resiliation anticipee. Chaque partie "
    "s engage a respecter les delais et a notifier tout incident dans les plus "
    "brefs delais. "
).split()


def make_text_pdf(path: Path, pages: int = 20, lines: int = 40) -> None:
    """Build a valid multi-page text PDF (correct xref/content references)."""
    content_per_page = []
    for page in range(pages):
        ops = ["BT /F1 11 Tf 50 760 Td"]
        for line in range(lines):
            text = f"Page {page + 1} line {line + 1}: " + " ".join(LOREM[:12]) + "."
            text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            if line:
                ops.append("0 -16 Td")
            ops.append(f"({text}) Tj")
        ops.append("ET")
        content_per_page.append("\n".join(ops).encode("latin-1"))

    objects = []
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    first_page_obj = 3
    kids = []
    page_obj_nums = []
    for i in range(pages):
        # page object: 3 + 2*i ; content object: 4 + 2*i
        kids.append(f"{first_page_obj + 2 * i} 0 R")
        page_obj_nums.append(first_page_obj + 2 * i)
    objects.append(f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {pages} >>".encode("latin-1"))
    font_obj = first_page_obj + 2 * pages
    for i in range(pages):
        content_obj = first_page_obj + 2 * i + 1
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_obj} 0 R >> >> "
            f"/Contents {content_obj} 0 R >>".encode("latin-1")
        )
        data = content_per_page[i]
        objects.append(b"<< /Length " + str(len(data)).encode() + b" >>\nstream\n" + data + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for idx, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{idx} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_pos = len(out)
    n = len(objects) + 1
    out += f"xref\n0 {n}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets[1:]:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {n} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


def make_docx_big(path: Path, paragraphs: int = 400) -> None:
    from docx import Document

    doc = Document()
    for i in range(paragraphs):
        doc.add_paragraph("Paragraphe %d: %s" % (i, " ".join(LOREM)))
    doc.save(path)


def make_xlsx_big(path: Path, rows: int = 800, cols: int = 8) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(rows):
        for c in range(cols):
            ws.cell(row=r + 1, column=c + 1, value=f"cell {r},{c} {LOREM[c % len(LOREM)]}")
    wb.save(path)


def make_odt_big(path: Path, paragraphs: int = 400) -> None:
    from odf.opendocument import OpenDocumentText
    from odf.text import P

    doc = OpenDocumentText()
    for i in range(paragraphs):
        doc.text.addElement(P(text=f"Paragraphe {i}: {' '.join(LOREM)}"))
    doc.save(str(path))


def make_ods_big(path: Path, rows: int = 800, cols: int = 6) -> None:
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    doc = OpenDocumentSpreadsheet()
    table = Table(name="Sheet1")
    doc.spreadsheet.addElement(table)
    for r in range(rows):
        row = TableRow()
        for c in range(cols):
            cell = TableCell()
            cell.addElement(P(text=f"r{r}c{c} {' '.join(LOREM)}"))
            row.addElement(cell)
        table.addElement(row)
    doc.save(str(path))


def build_fixtures(root: Path, n_per_type: int) -> list[Path]:
    root.mkdir(parents=True, exist_ok=True)
    big_text = "\n".join(" ".join(LOREM) for _ in range(300))
    for i in range(n_per_type):
        (root / f"t{i:03d}.txt").write_text(big_text, encoding="utf-8")
        (root / f"t{i:03d}.html").write_text(
            "<html><body>" + "".join(f"<p>{' '.join(LOREM)}</p>" for _ in range(200)) + "</body></html>",
            encoding="utf-8",
        )
        (root / f"t{i:03d}.json").write_text(
            json.dumps({"records": [{"i": j, "text": " ".join(LOREM)} for j in range(500)]}),
            encoding="utf-8",
        )
        make_text_pdf(root / f"t{i:03d}.pdf", pages=15)
        make_docx_big(root / f"t{i:03d}.docx", paragraphs=300)
        make_xlsx_big(root / f"t{i:03d}.xlsx", rows=600)
        make_odt_big(root / f"t{i:03d}.odt", paragraphs=300)
        make_ods_big(root / f"t{i:03d}.ods", rows=600)
    return sorted(p for p in root.iterdir() if p.is_file())


def run_sequential(paths: list[Path]) -> dict:
    from src.extractors.manager import ExtractionManager

    mgr = ExtractionManager(max_workers=1)
    ok = 0
    chars = 0
    per_type: dict[str, float] = {}
    with ResourceSampler() as s:
        for p in paths:
            t = time.perf_counter()
            r = mgr.extract_single(p)
            per_type[p.suffix] = per_type.get(p.suffix, 0.0) + (time.perf_counter() - t)
            ok += int(r.success)
            chars += len(r.content or "")
    return _summary("sequential", 1, len(paths), ok, chars, s, per_type)


def run_batch(paths: list[Path], workers: int) -> dict:
    from src.extractors.manager import ExtractionManager

    mgr = ExtractionManager(max_workers=workers)
    with ResourceSampler() as s:
        results = mgr.extract_batch(paths)
    ok = sum(1 for r in results.values() if r.success)
    chars = sum(len(r.content or "") for r in results.values())
    per_type: dict[str, float] = {}
    for p in paths:
        r = results.get(str(p))
        if r is not None:
            per_type[p.suffix] = per_type.get(p.suffix, 0.0) + r.extraction_time
    return _summary(f"batch_{workers}", workers, len(paths), ok, chars, s, per_type)


def _summary(kind, workers, n, ok, chars, s, per_type):
    wall = round(s.sample.wall, 3)
    return {
        "kind": kind,
        "workers": workers,
        "docs": n,
        "ok": ok,
        "chars": chars,
        "wall_s": wall,
        "docs_per_sec": round(n / wall, 1) if wall else 0,
        "cpu_pct": round(s.sample.cpu_pct, 1),
        "rss_mb": round(s.sample.rss_mb, 1),
        "iowait_pct": getattr(s, "iowait_pct", 0.0),
        "per_type_s": {k: round(v, 3) for k, v in sorted(per_type.items())},
    }


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--fixtures", default="/tmp/pis_bench/extract_real")
    ap.add_argument("--per-type", type=int, default=8)
    ap.add_argument("--workers", default="1,2,4")
    ap.add_argument("--regen", action="store_true")
    args = ap.parse_args()
    root = Path(args.fixtures)
    if args.regen or not root.exists():
        wipe(root)
        build_fixtures(root, args.per_type)
    paths = sorted(p for p in root.iterdir() if p.is_file())
    report = {"conditions": load_conditions(), "fixtures": len(paths),
              "bytes": sum(p.stat().st_size for p in paths), "results": []}
    # Warm imports + page cache on a representative subset, discard.
    run_sequential(paths[: min(len(paths), 8)])
    report["results"].append(run_sequential(paths))
    for w in [int(x) for x in args.workers.split(",") if x.strip()]:
        report["results"].append(run_batch(paths, w))
    json_print(report)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
