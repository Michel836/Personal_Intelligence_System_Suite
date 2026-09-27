"""Microsoft Office document extractor."""

import time
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from loguru import logger
from .base import BaseExtractor, ExtractionResult

try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    logger.warning("python-docx not available - DOCX extraction limited")

try:
    import openpyxl
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False
    logger.warning("openpyxl not available - Excel extraction disabled")


class OfficeExtractor(BaseExtractor):
    """Extract content from Microsoft Office documents."""
    
    def __init__(self):
        super().__init__()
        self.supported_extensions = {
            '.docx', '.doc',  # Word
            '.xlsx', '.xls',  # Excel  
            '.pptx', '.ppt',  # PowerPoint
        }
    
    def can_extract(self, file_path: Path) -> bool:
        """Check if file is a supported Office document."""
        return (
            file_path.suffix.lower() in self.supported_extensions and
            file_path.is_file()
        )
    
    def extract_content(self, file_path: Path) -> ExtractionResult:
        """Extract content from Office document."""
        start_time = time.time()
        extension = file_path.suffix.lower()
        
        try:
            if extension == '.docx':
                return self._extract_docx(file_path, start_time)
            elif extension == '.xlsx':
                return self._extract_xlsx(file_path, start_time)
            elif extension == '.pptx':
                return self._extract_pptx(file_path, start_time)
            elif extension in ['.doc', '.xls', '.ppt']:
                return ExtractionResult(
                    success=False,
                    error=f"Legacy format {extension} not supported (need modern Office formats)",
                    extraction_time=time.time() - start_time
                )
            else:
                return ExtractionResult(
                    success=False,
                    error=f"Unsupported extension: {extension}",
                    extraction_time=time.time() - start_time
                )
        
        except Exception as e:
            logger.debug(f"Error extracting Office document {file_path}: {e}")
            return ExtractionResult(
                success=False,
                error=f"Office extraction error: {str(e)}",
                extraction_time=time.time() - start_time
            )
    
    def _extract_docx(self, file_path: Path, start_time: float) -> ExtractionResult:
        """Extract text from DOCX file."""
        
        # Try with python-docx first (preferred)
        if DOCX_AVAILABLE:
            try:
                doc = DocxDocument(str(file_path))
                
                # Extract text from paragraphs
                text_content = []
                for paragraph in doc.paragraphs:
                    if paragraph.text.strip():
                        text_content.append(paragraph.text)
                
                # Extract text from tables
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            if cell.text.strip():
                                text_content.append(cell.text)
                
                full_text = '\n'.join(text_content)
                
                metadata = {
                    'paragraphs': len(doc.paragraphs),
                    'tables': len(doc.tables),
                    'characters': len(full_text),
                    'words': len(full_text.split()) if full_text else 0
                }
                
                # Try to get document properties
                try:
                    core_props = doc.core_properties
                    metadata.update({
                        'title': core_props.title or '',
                        'author': core_props.author or '',
                        'subject': core_props.subject or '',
                        'created': str(core_props.created) if core_props.created else '',
                        'modified': str(core_props.modified) if core_props.modified else ''
                    })
                except:
                    pass
                
                return ExtractionResult(
                    success=True,
                    content=full_text,
                    metadata=metadata,
                    extraction_time=time.time() - start_time
                )
            
            except Exception as e:
                logger.debug(f"python-docx failed for {file_path}, trying manual extraction: {e}")
        
        # Fallback: manual XML extraction
        return self._extract_docx_manual(file_path, start_time)
    
    def _extract_docx_manual(self, file_path: Path, start_time: float) -> ExtractionResult:
        """Manual DOCX extraction via ZIP parsing."""
        try:
            text_content = []
            
            with zipfile.ZipFile(file_path, 'r') as zip_file:
                # Extract main document
                try:
                    with zip_file.open('word/document.xml') as xml_file:
                        tree = ET.parse(xml_file)
                        root = tree.getroot()
                        
                        # Extract text from all text nodes
                        for text_elem in root.iter():
                            if text_elem.tag.endswith('}t') and text_elem.text:
                                text_content.append(text_elem.text)
                
                except KeyError:
                    pass
            
            full_text = ' '.join(text_content)
            
            metadata = {
                'extraction_method': 'manual_xml',
                'characters': len(full_text),
                'words': len(full_text.split()) if full_text else 0
            }
            
            return ExtractionResult(
                success=bool(full_text.strip()),
                content=full_text,
                metadata=metadata,
                extraction_time=time.time() - start_time
            )
        
        except Exception as e:
            return ExtractionResult(
                success=False,
                error=f"Manual DOCX extraction failed: {str(e)}",
                extraction_time=time.time() - start_time
            )
    
    def _extract_xlsx(self, file_path: Path, start_time: float) -> ExtractionResult:
        """Extract text from Excel XLSX file."""
        
        if not EXCEL_AVAILABLE:
            return ExtractionResult(
                success=False,
                error="openpyxl not available",
                extraction_time=time.time() - start_time
            )
        
        try:
            workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            
            text_content = []
            total_cells = 0
            
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                
                # Add sheet name
                text_content.append(f"Sheet: {sheet_name}")
                
                # Extract cell values (limit to prevent memory issues)
                cell_count = 0
                for row in sheet.iter_rows(max_row=1000, max_col=50):  # Limit range
                    row_text = []
                    for cell in row:
                        if cell.value is not None:
                            row_text.append(str(cell.value))
                            cell_count += 1
                    
                    if row_text:
                        text_content.append('\t'.join(row_text))
                    
                    if cell_count > 10000:  # Stop if too many cells
                        break
                
                total_cells += cell_count
            
            workbook.close()
            
            full_text = '\n'.join(text_content)
            
            metadata = {
                'sheets': len(workbook.sheetnames),
                'cells_processed': total_cells,
                'characters': len(full_text),
                'words': len(full_text.split()) if full_text else 0
            }
            
            return ExtractionResult(
                success=True,
                content=full_text,
                metadata=metadata,
                extraction_time=time.time() - start_time
            )
        
        except Exception as e:
            return ExtractionResult(
                success=False,
                error=f"Excel extraction error: {str(e)}",
                extraction_time=time.time() - start_time
            )
    
    def _extract_pptx(self, file_path: Path, start_time: float) -> ExtractionResult:
        """Extract text from PowerPoint PPTX file."""
        
        try:
            text_content = []
            
            with zipfile.ZipFile(file_path, 'r') as zip_file:
                # Get list of slide files
                slide_files = [f for f in zip_file.namelist() if f.startswith('ppt/slides/slide') and f.endswith('.xml')]
                
                for slide_file in sorted(slide_files):
                    try:
                        with zip_file.open(slide_file) as xml_file:
                            tree = ET.parse(xml_file)
                            root = tree.getroot()
                            
                            # Extract text from all text nodes
                            slide_text = []
                            for text_elem in root.iter():
                                if text_elem.tag.endswith('}t') and text_elem.text:
                                    slide_text.append(text_elem.text)
                            
                            if slide_text:
                                text_content.append('Slide: ' + ' '.join(slide_text))
                    
                    except Exception as e:
                        logger.debug(f"Error extracting slide {slide_file}: {e}")
                        continue
            
            full_text = '\n'.join(text_content)
            
            metadata = {
                'slides': len(slide_files),
                'extraction_method': 'xml_parsing',
                'characters': len(full_text),
                'words': len(full_text.split()) if full_text else 0
            }
            
            return ExtractionResult(
                success=bool(full_text.strip()),
                content=full_text,
                metadata=metadata,
                extraction_time=time.time() - start_time
            )
        
        except Exception as e:
            return ExtractionResult(
                success=False,
                error=f"PowerPoint extraction error: {str(e)}",
                extraction_time=time.time() - start_time
            )