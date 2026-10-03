"""Lite acceptance: real UI buttons, real extraction, temporary source corpus."""
import threading
from pathlib import Path

import pytest
from streamlit.testing.v1 import AppTest

from src.core.database import DatabaseManager
from src.core.lite_workflow import LiteJob
from src.core.scan_service import ScanRequest, ScanService
from tests.release.corpus import _minimal_pdf

APP = str(Path(__file__).resolve().parents[2] / "src/ui/app.py")


@pytest.fixture
def lite(tmp_path, monkeypatch):
    monkeypatch.setenv("PIS_LAUNCH_PROFILE", "lite")
    monkeypatch.setenv("PIS_DB_PATH", str(tmp_path / "index.db"))
    monkeypatch.setenv("PIS_OCR_ENABLED", "0")
    monkeypatch.setenv("PIS_BACKUP_DIR", str(tmp_path / "backups"))
    root = tmp_path / "corpus"
    root.mkdir()
    (root / "résumé.txt").write_text("candidat immunothérapie unicorne", encoding="utf-8")
    (root / "notes.md").write_text("# RAG\narchitecte citrouille", encoding="utf-8")
    (root / "report.pdf").write_bytes(_minimal_pdf("finance papillon"))
    from docx import Document
    doc = Document()
    doc.add_paragraph("contrat framboise")
    doc.save(root / "contrat.docx")
    from openpyxl import Workbook
    book = Workbook()
    book.active["A1"] = "agenda noisette"
    book.save(root / "agenda.xlsx")
    (root / "broken.pdf").write_bytes(b"invalid PDF")
    (root / "empty.txt").touch()
    (root / "photo.jpg").write_bytes(b"not an image")
    return root


def app():
    at = AppTest.from_file(APP, default_timeout=30).run()
    assert not at.exception, [e.message for e in at.exception]
    return at


def button(at, label):
    return next(b for b in at.button if b.label == label)


def nav(at, label):
    at.radio[0].set_value(label).run()
    assert not at.exception
    return at


def finish(at):
    job = at.session_state["lite_job"]
    job.join(30)
    assert not job.running
    at.run()
    assert not at.exception, [e.message for e in at.exception]
    return job.snapshot()


def test_real_ui_scan_extract_search_restart_backup(lite, tmp_path):
    before = {p.name: p.read_bytes() for p in lite.iterdir()}
    at = app()
    assert len(at.radio[0].options) == 4
    at.text_input[0].set_value(str(lite))
    button(at, "Lancer le scan").click().run()
    assert finish(at)["status"] == "COMPLETED"
    nav(at, "🔄 Extraire")
    button(at, "Extraire le prochain lot").click().run()
    state = finish(at)
    assert state["status"] == "COMPLETED"
    assert state["counts"]["EXTRACTED"] == 5
    assert sum(state["counts"].values()) == 6  # malformed PDF recorded once
    nav(at, "🔍 Search")
    for word, name in [("unicorne", "résumé.txt"), ("citrouille", "notes.md"),
                       ("papillon", "report.pdf"), ("framboise", "contrat.docx"),
                       ("noisette", "agenda.xlsx")]:
        at.text_input[0].set_value(word)
        button(at, "🔍 Search").click().run()
        assert not at.exception
        assert name in [e.label for e in at.expander]
    restarted = nav(app(), "🔍 Search")
    restarted.text_input[0].set_value("unicorne")
    button(restarted, "🔍 Search").click().run()
    assert "résumé.txt" in [e.label for e in restarted.expander]
    nav(at, "🧰 État du système")
    at.text_input[0].set_value(str(tmp_path / "backups"))
    button(at, "Sauvegarder l’index").click().run()
    assert not at.exception
    assert at.success
    from src.ops.backup import verify_backup
    archives = list((tmp_path / "backups").glob("*.tar.gz"))
    assert len(archives) == 1
    assert verify_backup(archives[0])["ok"]
    assert before == {p.name: p.read_bytes() for p in lite.iterdir()}


def test_invalid_empty_and_malformed_search(lite, tmp_path):
    at = app()
    at.text_input[0].set_value(str(tmp_path / "absent"))
    button(at, "Lancer le scan").click().run()
    assert at.error and not at.exception
    empty = tmp_path / "empty"
    empty.mkdir()
    at.text_input[0].set_value(str(empty))
    button(at, "Lancer le scan").click().run()
    assert finish(at)["processed"] == 0
    nav(at, "🔍 Search")
    for query in ["---", 'foo"bar', "AND", "NEAR("]:
        at.text_input[0].set_value(query)
        button(at, "🔍 Search").click().run()
        assert not at.exception


def test_incremental_modify_delete_and_cancel(lite, tmp_path):
    db = DatabaseManager(tmp_path / "index.db")
    job = LiteJob(db.db_path)
    job.start("scan", str(lite))
    job.join()
    assert job.snapshot()["status"] == "COMPLETED"
    original = db.get_stats()["total_files"]
    job.start("scan", str(lite))
    job.join()
    assert db.get_stats()["total_files"] == original
    job.start("extract", str(lite))
    job.join()
    path = lite / "résumé.txt"
    path.write_text("nouveau contenu tangerine", encoding="utf-8")
    job.start("scan", str(lite))
    job.join()
    job.start("extract", str(lite))
    job.join()
    assert db.search_files(query="tangerine")
    assert not db.search_files(query="unicorne")
    path.unlink()
    cancel = threading.Event()
    cancel.set()
    result = ScanService(db).run(ScanRequest(lite), cancel_event=cancel)
    assert result.status == "CANCELLED"
    assert db.search_files(query="tangerine")  # cancelled scan never reconciles
    job.start("scan", str(lite))
    job.join()
    assert not db.search_files(query="tangerine")
    assert db.search_files(query="tangerine", include_missing=True)


def test_producer_exception_is_failed_without_reconciliation(lite, tmp_path, monkeypatch):
    from src.scanner.fast_engine import FastScannerEngine
    db = DatabaseManager(tmp_path / "index.db")
    assert ScanService(db).run(ScanRequest(lite)).status == "COMPLETED"
    count = db.get_stats()["total_files"]
    def broken(self, *args, **kwargs):
        raise PermissionError("unreadable subtree")
        yield
    monkeypatch.setattr(FastScannerEngine, "scan_paths", broken)
    result = ScanService(db).run(ScanRequest(lite))
    assert result.status == "FAILED"
    assert "unreadable subtree" in result.error
    assert db.get_stats()["total_files"] == count


def test_ui_stop_targets_worker_and_new_tab_shares_job(lite, monkeypatch):
    from src.scanner.fast_engine import FastScannerEngine
    entered = threading.Event()
    release = threading.Event()
    original = FastScannerEngine.scan_paths
    def controlled(self, *args, **kwargs):
        entered.set()
        assert release.wait(10)
        yield from original(self, *args, **kwargs)
    monkeypatch.setattr(FastScannerEngine, "scan_paths", controlled)
    at = app()
    at.text_input[0].set_value(str(lite))
    button(at, "Lancer le scan").click().run()
    assert entered.wait(5)
    try:
        second = app()
        assert second.session_state["lite_job"] is at.session_state["lite_job"]
        assert button(second, "Lancer le scan").disabled
        button(at, "Arrêter").click().run()
        assert at.session_state["lite_job"].cancel_event.is_set()
    finally:
        release.set()
    assert finish(at)["status"] == "CANCELLED"


def test_traversal_error_never_reports_success(lite, tmp_path, monkeypatch):
    import src.scanner.fast_engine as engine
    def denied(*args, onerror=None, **kwargs):
        onerror(PermissionError("subtree denied"))
        yield
    monkeypatch.setattr(engine.os, "walk", denied)
    result = ScanService(DatabaseManager(tmp_path / "index.db")).run(ScanRequest(lite))
    assert result.status == "FAILED"
    assert "subtree denied" in result.error


def test_stale_scan_recovered_only_on_first_app_start(lite, tmp_path):
    db = DatabaseManager(tmp_path / "index.db")
    run = db.begin_scan(lite)
    at = app()
    with db.get_connection() as conn:
        assert conn.execute("SELECT status FROM scan_runs WHERE id=?", (run["run_id"],)).fetchone()[0] == "FAILED"
    at.text_input[0].set_value(str(lite))
    button(at, "Lancer le scan").click().run()
    assert finish(at)["status"] == "COMPLETED"


def test_file_access_error_preserves_existing_index(lite, tmp_path, monkeypatch):
    db = DatabaseManager(tmp_path / "index.db")
    service = ScanService(db)
    assert service.run(ScanRequest(lite)).status == "COMPLETED"
    original_stat = Path.stat
    def denied(path, *args, **kwargs):
        if path == lite / "agenda.xlsx":
            raise PermissionError("file denied")
        return original_stat(path, *args, **kwargs)
    monkeypatch.setattr(Path, "stat", denied)
    result = service.run(ScanRequest(lite))
    assert result.status == "FAILED"
    assert result.files_errors == 1
    assert db.search_files(query="agenda")[0]["state"] == "ACTIVE"
